"""Generate UIEP-3745 demo Word document with steps and screenshots."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "Demo-Screenshots"
FIGMA = ROOT / "figma screenshot"
OUT = ROOT / "UIEP-3745_COE_Demo_Document.docx"


def set_run_font(run, size=11, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return h


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def add_image(doc, path: Path, width_inches=6.0, caption=None):
    if not path.exists():
        add_para(doc, f"[Screenshot not found: {path.name}]", bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        set_run_font(r, size=9, bold=False)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_step(doc, title, goal, account, steps, expected, image_path=None, figma_path=None, image_caption=None):
    add_heading(doc, title, level=2)
    add_para(doc, f"Goal: {goal}", bold=True)
    if account:
        add_para(doc, f"Account: {account}")
    add_para(doc, "Reproduce steps:", bold=True)
    for i, s in enumerate(steps, 1):
        p = doc.add_paragraph(f"{i}. {s}")
        for run in p.runs:
            set_run_font(run)
    add_para(doc, "Expected result:", bold=True)
    add_bullets(doc, expected)
    add_para(doc, "Pass / Fail:  ☐ Pass    ☐ Fail", bold=True)
    if image_path:
        add_para(doc, "Screen reference (actual / demo capture):", bold=True)
        add_image(doc, image_path, width_inches=6.2, caption=image_caption or image_path.name)
    if figma_path and figma_path.exists():
        add_para(doc, "Figma reference:", bold=True)
        add_image(doc, figma_path, width_inches=5.8, caption=f"Figma — {figma_path.name}")
    doc.add_paragraph()


def main():
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("UIEP-3745 — Corporate Officer Exclusion")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Demo / UAT Document — Reproducible Steps with Screens")
    set_run_font(r, size=14, bold=True)

    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Ticket", "UIEP-3745"],
            ["Feature", "COE election end messaging & confirmation flow"],
            ["Route", "/manage-account/corporate-office-exclusion"],
            ["Release", "Oct 28 UAT / NNMilestone2"],
            ["Demo date", "____________________"],
            ["Tester", "____________________"],
            ["Environment", "Local / UAT ____________________"],
        ],
    )

    add_heading(doc, "1. Purpose", level=1)
    add_para(
        doc,
        "Demonstrate that Corporate Officer Exclusion no longer fails silently after Submit, "
        "shows clear success/error messaging, blocks a second election attempt on revisit, "
        "and matches Figma: Intro → Worksheet → Success → Already elected.",
    )

    add_heading(doc, "2. Screen map", level=1)
    add_table(
        doc,
        ["#", "Screen", "Capture / Figma"],
        [
            ["A", "Before fix (baseline)", "00-before-initial-worksheet.png"],
            ["B", "Intro", "06-intro-screen.png + Initial.png"],
            ["C", "Must-calculate banner", "01-submit-without-calculate-must-calculate-banner.png"],
            ["D", "Inline missing value", "Figma: inline errors.png"],
            ["E", "No-savings warning", "02-no-savings-warning-on-submit.png"],
            ["F", "Eligibility reject", "03-eligibility-rejected-deadline-and-payroll.png"],
            ["G", "Success", "Figma: Success.png (needs eligible account)"],
            ["H", "Already elected", "Figma: Already visited.png (needs elected account)"],
        ],
    )

    add_heading(doc, "3. Accounts required", level=1)
    add_table(
        doc,
        ["Scenario", "Account type"],
        [
            ["Happy path (Intro → Success)", "Eligible newly subject Corp / LLC-as-Corp (<$500k, in window)"],
            ["Already elected", "Account with COE already in effect"],
            ["Negative paths", "Any COE-openable account (e.g. VANDERLOOP)"],
        ],
    )
    add_para(
        doc,
        "Eligible tip (from ticket): Register a Corporation that activates on submit, then take COE "
        "before initial report due date OR within 30 days of Initial Determination.",
    )

    add_heading(doc, "4. Pre-demo checklist", level=1)
    add_bullets(
        doc,
        [
            "Office build succeeds (0 errors) with latest COE files",
            "Logged in with access to test employer account",
            "Navigate: View & Manage Account → Corporate Officer Exclusion",
            "Have this Word doc open for screen compare",
        ],
    )

    add_heading(doc, "5. Reproducible demo steps", level=1)

    add_step(
        doc,
        "STEP 0 — Baseline (what was broken)",
        "Show original defect: Submit did nothing",
        "Any (old behavior / capture)",
        [
            "Open baseline capture or old build",
            "On worksheet click SUBMIT ELECTION with no end messaging",
        ],
        [
            "Before: no success screen, no error banner (silent)",
            "Ticket: “Nothing happens after successful COE election submission”",
        ],
        image_path=SHOTS / "00-before-initial-worksheet.png",
        image_caption="Baseline — worksheet only; Submit appeared to do nothing",
    )

    add_step(
        doc,
        "STEP 1 — Intro screen",
        "Landing shows Intro (not worksheet) with eligibility + NEXT",
        "Not-yet-elected account (e.g. 182138-000-5)",
        [
            "Select employer account",
            "Open Corporate Officer Exclusion",
            "Do not click NEXT yet — capture/show Intro",
        ],
        [
            "Title: Corporate Officer Exclusion",
            "Eligibility copy ($500k / Corp / LLC-as-Corp)",
            "Grey card: Account Number (######-###-#), Legal Name, effective-as-of date",
            "NEXT button bottom-right",
        ],
        image_path=SHOTS / "06-intro-screen.png",
        figma_path=FIGMA / "Initial.png",
        image_caption="Actual Intro screen (demo capture)",
    )

    add_step(
        doc,
        "STEP 2 — Worksheet after NEXT",
        "NEXT opens reasonableness worksheet",
        "Same as Step 1",
        ["From Intro click NEXT"],
        [
            "Title: Corporate Officer Exclusion Reasonableness Worksheet",
            "Table lines 1–10 + Net Savings",
            "Buttons: CANCEL | CALCULATE | SUBMIT ELECTION",
        ],
        figma_path=FIGMA / "Officer Exclusion Calculate Worksheet.png",
    )

    add_step(
        doc,
        "STEP 3 — Submit without Calculate",
        "Submit is never silent",
        "Same worksheet",
        [
            "On worksheet click SUBMIT ELECTION without clicking CALCULATE",
        ],
        [
            "Red Alert: “You must calculate your net savings before electing this exclusion.”",
            "Stay on worksheet",
        ],
        image_path=SHOTS / "01-submit-without-calculate-must-calculate-banner.png",
        image_caption="Must-calculate banner after Submit without Calculate",
    )

    add_step(
        doc,
        "STEP 4 — Inline missing value",
        "Required $0 fields show Missing value",
        "Same worksheet",
        [
            "Leave Line 1 FUTA Without at $0.00 (and/or WI Without at $0)",
            "Click CALCULATE",
        ],
        [
            "Top Alert lists missing fields (e.g. Payroll Subject to FUTA missing.)",
            "Field shows red border + “Missing value”",
        ],
        figma_path=FIGMA / "Officer Exclusion Calculate Worksheet inline errors.png",
    )

    add_step(
        doc,
        "STEP 5 — Calculate updates values",
        "CALCULATE refreshes tax rows / Net Savings",
        "Same worksheet",
        [
            "Enter FUTA Without > 0 (e.g. 400)",
            "Enter WI Without > 0 (e.g. 600)",
            "Enter WI With (e.g. 0)",
            "Click CALCULATE",
        ],
        [
            "CALCULATING… then values update",
            "Error banners clear",
        ],
        figma_path=FIGMA / "Officer Exclusion Caluclator Worksheet.png",
    )

    add_step(
        doc,
        "STEP 6A — Submit success (happy path) ★ eligible account",
        "End messaging after successful election",
        "Eligible newly subject Corp (required)",
        [
            "Complete Step 5",
            "Click SUBMIT ELECTION",
            "If no-savings warning (6B), click SUBMIT ELECTION again",
            "Observe success screen",
        ],
        [
            "Green: “Your election successfully submitted.”",
            "Links: View Initial Determination letter; Mail a copy…",
            "Summary card: Account, Legal Name, Confirmation ID (if API), Date Submitted (if API), effective date",
            "ACCOUNT SUMMARY only (no BACK)",
        ],
        figma_path=FIGMA / "Officer Exclusion Calculator -Success.png",
    )

    add_step(
        doc,
        "STEP 6B — No-savings / pay MORE warning (optional)",
        "Backend warning when estimates show higher tax",
        "Account that can submit with negative net savings",
        [
            "Calculate with figures that produce negative Net Savings",
            "Click SUBMIT ELECTION once",
        ],
        [
            "Alert warns you will pay MORE unemployment tax…",
            "Can Submit again to proceed or Cancel",
        ],
        image_path=SHOTS / "02-no-savings-warning-on-submit.png",
        image_caption="No-savings / pay MORE warning",
    )

    add_step(
        doc,
        "STEP 6C — Mail copy banner",
        "Mail-copy link shows Figma banner",
        "After successful submit",
        ["On success click “Mail a copy of the ID to the main business mailing address”"],
        [
            "Banner: “Initial Determination letter will be mailed to your main business mailing address.”",
        ],
        figma_path=FIGMA / "Officer Exclusion Calculator -Success with banner.png",
    )

    add_step(
        doc,
        "STEP 6D — Account Summary exit",
        "End button returns to Account Summary",
        "Success screen",
        ["Click ACCOUNT SUMMARY"],
        [
            "Navigates to /manage-account/account-summary",
            "Not account-details; not legacy Complete",
        ],
    )

    add_step(
        doc,
        "STEP 7 — Already elected (revisit) ★ elected account",
        "Cannot take election again; confirmation shown",
        "Same account after Step 6A success, or known elected account",
        [
            "From Account Summary open Corporate Officer Exclusion again",
        ],
        [
            "No worksheet / no Submit",
            "Banner: “Corporate Officer Exclusion submitted.”",
            "Card with account details + ACCOUNT SUMMARY only",
        ],
        figma_path=FIGMA / "Officer Exclusion Calculator -Already visited.png",
    )

    add_step(
        doc,
        "STEP 8 — Eligibility reject (optional)",
        "Backend rule violations display clearly",
        "Ineligible (deadline / ≥$500k) e.g. VANDERLOOP",
        ["Intro → NEXT → Calculate → Submit"],
        [
            "Alert shows deadline and/or $500,000 payroll messages",
        ],
        image_path=SHOTS / "03-eligibility-rejected-deadline-and-payroll.png",
        image_caption="Eligibility rejected — deadline and payroll rules",
    )

    add_step(
        doc,
        "STEP 9 — Submit failed banner (if forceable)",
        "Unexpected save failure messaging",
        "Dev-only / N/A in UAT if not practical",
        ["Force save exception if possible; otherwise mark N/A"],
        ["Banner: “Worksheet failed to submit. Please try again.”"],
        figma_path=FIGMA / "Officer Exclusion Calculate Worksheet submit failed.png",
    )

    add_heading(doc, "6. Ticket acceptance mapping", level=1)
    add_table(
        doc,
        ["Ticket note", "Demo step"],
        [
            ["Nothing after successful submit", "Step 6A"],
            ["Can Submit again on same page", "Steps 6A + 7"],
            ["No COE effective date", "Steps 1, 6A, 7"],
            ["No Account Summary navigation", "Step 6D"],
            ["No intro screen", "Step 1"],
            ["Already-elected messaging", "Step 7"],
            ["Account Summary (not Complete)", "Step 6D"],
            ["View / mail Initial Determination", "Steps 6A + 6C (UI)"],
        ],
    )

    add_heading(doc, "7. Out of scope", level=1)
    add_bullets(
        doc,
        [
            "Dual already-in-effect + suspended errors → UIEP-3750 (backend)",
            "Confirmation ID exact contract vs CorporateOfficerExclusionSK → Mendota/product",
            "Actual mail send of ID letter → backend/product",
        ],
    )

    add_heading(doc, "8. Demo talking points (2 minutes)", level=1)
    add_bullets(
        doc,
        [
            "Before: Submit did nothing → employer could retry and get confusing errors",
            "Now: Intro → Calculate required → clear banners",
            "Success: green confirmation + letter links + Account Summary",
            "Revisit: already-elected card — no second Submit on worksheet",
        ],
    )

    add_heading(doc, "9. Sign-off", level=1)
    add_table(
        doc,
        ["Role", "Name", "Date", "Result"],
        [
            ["Developer", "", "", "☐ Pass  ☐ Fail"],
            ["QA / UAT", "", "", "☐ Pass  ☐ Fail"],
            ["Product (optional)", "", "", "☐ Pass  ☐ Fail"],
        ],
    )
    add_para(doc, "Notes / defects found during demo:")
    add_para(doc, "_" * 80)
    add_para(doc, "_" * 80)

    doc.save(OUT)
    print(f"Wrote: {OUT}")


if __name__ == "__main__":
    main()
